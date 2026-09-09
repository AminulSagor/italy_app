from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = Path("docs/Italir_Pothe_App_Store_Purchase_Mapping_Guide.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

INK = "17231B"
GREEN = "006B3F"
PALE_GREEN = "EAF4EE"
BLUE = "173B57"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F4F6F5"
MID_GRAY = "66736A"
GRID = "D9D9D9"
AMBER = "8A5A00"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(table, color=GRID, size="6"):
    tbl_pr = table._tbl.tblPr
    old = tbl_pr.find(qn("w:tblBorders"))
    if old is not None:
        tbl_pr.remove(old)
    node = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        node.append(el)
    tbl_pr.append(node)


def margins(cell, top=90, start=110, bottom=90, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = tc_mar.find(qn(f"w:{name}"))
        if el is None:
            el = OxmlElement(f"w:{name}")
            tc_mar.append(el)
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    rep = OxmlElement("w:tblHeader")
    rep.set(qn("w:val"), "true")
    tr_pr.append(rep)


def set_cell_text(cell, text, *, bold=False, color=INK, size=9.2, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Aptos"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    margins(cell)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.extend([color, underline])
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(9)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def title(doc, text, subtitle=None, page_break=False):
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_after = Pt(10)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if page_break:
        p.add_run().add_break(WD_BREAK.PAGE)
    r = p.add_run(text)
    r.font.color.rgb = RGBColor.from_string("000000")
    if subtitle:
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(16)
        r2 = p2.add_run(subtitle)
        r2.font.size = Pt(14)
        r2.font.color.rgb = RGBColor.from_string(MID_GRAY)


def h(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.color.rgb = RGBColor.from_string("000000")
    return p


def para(doc, text="", *, bold_lead=None, italic=False, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.12
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        r1.bold = True
        r1.font.color.rgb = RGBColor.from_string(INK)
        r2 = p.add_run(text[len(bold_lead):])
        r2.font.color.rgb = RGBColor.from_string(INK)
    else:
        r = p.add_run(text)
        r.italic = italic
        r.font.color.rgb = RGBColor.from_string(INK)
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.08
    p.add_run(text)
    return p


def step(doc, n, head, body):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.08)
    p.paragraph_format.first_line_indent = Inches(-0.08)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(f"{n}. {head} ")
    r.bold = True
    p.add_run(body)


def table(doc, headers, rows, widths=None, font_size=9.0):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    borders(t)
    set_repeat_header(t.rows[0])
    for i, value in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], value, bold=True, color="FFFFFF", size=9.0, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(t.rows[0].cells[i], BLUE)
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], str(value), size=font_size, align=WD_ALIGN_PARAGRAPH.LEFT)
            if ridx % 2:
                shade(cells[i], PALE_BLUE)
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.68)
sec.bottom_margin = Inches(0.66)
sec.left_margin = Inches(0.72)
sec.right_margin = Inches(0.72)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(11)
styles["Normal"].font.color.rgb = RGBColor.from_string(INK)
styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(30)
styles["Title"].font.bold = True
title_ppr = styles["Title"]._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)
for i, size in ((1, 19), (2, 14), (3, 11.5)):
    st = styles[f"Heading {i}"]
    st.font.name = "Aptos Display"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string("000000")

for section in doc.sections:
    hp = section.header.paragraphs[0]
    hp.text = "ITALIR POTHE   APP STORE COMMERCE"
    hp.runs[0].font.name = "Aptos"
    hp.runs[0].font.size = Pt(8.5)
    hp.runs[0].font.bold = True
    hp.runs[0].font.color.rgb = RGBColor.from_string(MID_GRAY)
    add_page_number(section.footer.paragraphs[0])

title(doc, "App Store Purchase Mapping Guide", "Courses, AI credits, CV credits, streak protection, subscriptions, and coupons")
para(doc, "Product and backend configuration guide for Italir Pothe", bold_lead="Product and backend configuration guide")
para(doc, "Bundle ID  com.shafacode.italirpothe")
para(doc, "Prepared  7 September 2026")

doc.add_paragraph().paragraph_format.space_after = Pt(14)
h(doc, "Purpose and decision", 1)
para(doc, "This guide defines which Apple product type to create for every paid feature in the current Italir Pothe mobile and backend code, how to map that product in the admin portal, and how coupons must be implemented. The decisive rule is that Apple’s Product ID, the admin Product ID, and the backend transaction Product ID must match exactly.")
para(doc, "Do not use the backend EUR amount as the customer charge. StoreKit returns the localized App Store price, and Apple controls tax, currency conversion, and storefront pricing.")

h(doc, "Current implementation boundary", 2)
bullet(doc, "Monthly auto-renewal is supported end to end only for unlimited streak protection.")
bullet(doc, "Limited-day course access is a non-renewing subscription, not an auto-renewable subscription.")
bullet(doc, "AI and CV monthly package fields exist in parts of the schema and UI, but the iOS subscription lifecycle rejects anything except monthly unlimited streak protection. Keep AI and CV packages one-time until the lifecycle is extended.")
bullet(doc, "The exact products already present in App Store Connect cannot be verified from source code. Use the screenshot checklist near the end to complete an account-level reconciliation.")

h(doc, "Contents", 2)
for item in [
    "1  Master product mapping",
    "2  Product ID and metadata rules",
    "3  App Store Connect setup procedures",
    "4  Admin portal mapping",
    "5  Coupon architecture",
    "6  Server configuration and notifications",
    "7  Test and review checklist",
    "8  Screenshot reconciliation worksheet",
]:
    para(doc, item, space_after=2)

title(doc, "Master Product Mapping", page_break=True)
para(doc, "Create the Apple product first, then copy the immutable Product ID into the matching Italir Pothe admin record.")

mapping_rows = [
    ("Lifetime course", "Non Consumable", "One purchase unlocks one course permanently", "Course mapping  lifetime  non_consumable", "Restorable"),
    ("Limited day course", "Non Renewing Subscription", "Access expires after the configured number of days and does not auto renew", "Course mapping  time_limited  subscription", "Backend managed expiry"),
    ("AI credit bundle", "Consumable", "Voice minutes and text tokens are depleted by use", "Package  ai_bundle  one_time  consumable", "Not restored by Apple"),
    ("CV credit pack", "Consumable", "Each credit is spent on a new paid CV creation", "Package  cv_credit  one_time  consumable", "Not restored by Apple"),
    ("Finite streak freeze pack", "Consumable", "A purchased quantity is depleted when missed days are protected", "Package  streak_freeze  finite  one_time  consumable", "Not restored by Apple"),
    ("Monthly unlimited streak protection", "Auto Renewable Subscription", "Protection remains active while the monthly subscription is valid", "Package  streak_freeze  monthly_unlimited  monthly  subscription", "Restorable and renewable"),
]
table(doc, ["Feature", "Apple type", "Entitlement behavior", "Admin and backend values", "Recovery"], mapping_rows, [1.25, 1.25, 2.0, 1.75, 1.0], 8.35)

h(doc, "Do not create these mappings", 2)
bullet(doc, "Do not map AI credits, CV credits, or finite streak freezes as non-consumables. The backend requires consumable products for one-time packages.")
bullet(doc, "Do not map a lifetime course as consumable. The backend requires a non-consumable transaction type.")
bullet(doc, "Do not put a limited-day course in the auto-renewable Streak Protection subscription group. It is a non-renewing subscription and has its own expiry date.")
bullet(doc, "Do not create a monthly AI or monthly CV subscription for production with the current backend. Renewal registration supports only monthly unlimited streak protection.")

h(doc, "Recommended Product ID pattern", 2)
para(doc, "Product IDs are immutable and cannot be reused after deletion. Use lowercase, stable semantic names and a version suffix only when the commercial meaning changes.")
table(doc, ["Feature", "Recommended example"], [
    ("Lifetime course", "com.shafacode.italirpothe.course.a1.lifetime.v1"),
    ("Limited course", "com.shafacode.italirpothe.course.a1.90day.v1"),
    ("AI bundle", "com.shafacode.italirpothe.ai.bundle.small.v1"),
    ("CV credits", "com.shafacode.italirpothe.cv.credits.5.v1"),
    ("Finite freezes", "com.shafacode.italirpothe.streak.freezes.5.v1"),
    ("Monthly streak", "com.shafacode.italirpothe.streak.unlimited.monthly.v1"),
], [2.1, 5.1], 9.4)
title(doc, "Product Metadata Rules", page_break=True)
h(doc, "Fields shared by in app purchases", 1)
table(doc, ["App Store Connect field", "How to complete it"], [
    ("Reference Name", "Internal label, maximum 64 characters. Example: CV Credits 5 v1."),
    ("Product ID", "Exact immutable identifier used by StoreKit and backend. Apple permits letters, numbers, hyphens, periods, and underscores, up to 100 characters."),
    ("Price and Availability", "Choose the App Store price point and storefront availability. This is the real customer charge."),
    ("Display Name", "Customer-facing localized name, 2 to 30 characters."),
    ("Description", "Customer-facing localized description, maximum 45 characters."),
    ("App Review Screenshot", "Show the in-app paywall or purchase screen and the exact product or service. It is review-only."),
    ("Review Notes", "Explain navigation, entitlement, expiry or depletion, restore behavior, and provide a review account if sign-in is required."),
], [1.65, 5.55], 9.35)

h(doc, "Suggested customer metadata", 2)
table(doc, ["Product", "Display Name", "Description"], [
    ("Lifetime course", "Italian A1 Course", "Lifetime access to the complete A1 course"),
    ("90 day course", "A1 Course 90 Days", "Access to the A1 course for 90 days"),
    ("AI bundle", "AI Tutor Credit Pack", "More AI voice minutes and text credits"),
    ("CV pack", "5 CV Credits", "Create five additional professional CVs"),
    ("Finite freezes", "5 Streak Freezes", "Protect five missed learning days"),
    ("Monthly streak", "Monthly Streak Protection", "Unlimited missed-day protection each month"),
], [1.45, 2.1, 3.65], 9.15)
para(doc, "Confirm every localized string remains within Apple’s field limits before saving. The final wording should match the app screen shown in the review screenshot.")

h(doc, "Review notes template", 2)
para(doc, "Navigate to Menu, then Shop, then select the relevant category. The app retrieves the localized price from StoreKit. After Apple confirms the transaction, our server verifies the signed transaction and grants the displayed entitlement. Consumable balances decrease when used. Lifetime courses can be restored. Monthly streak protection renews automatically and can be managed in Apple subscription settings. Review account: insert credentials and any required test path.")

title(doc, "App Store Connect Setup Procedures", page_break=True)
h(doc, "One time consumables", 1)
para(doc, "Use this procedure for AI bundles, CV credit packs, and finite streak-freeze packs.")
step(doc, 1, "Open the app", "In App Store Connect, open Apps, select Italir Pothe, then open Monetization and In-App Purchases.")
step(doc, 2, "Create the product", "Click the add control and select Consumable.")
step(doc, 3, "Enter identity", "Add the Reference Name and the final Product ID. Recheck the ID before creating because it cannot be edited or reused.")
step(doc, 4, "Set commerce data", "Choose price, tax category if applicable, and country or region availability.")
step(doc, 5, "Add localization", "Enter a customer-facing Display Name and Description for at least one language.")
step(doc, 6, "Add review material", "Upload the in-app purchase screenshot and add concise review notes describing where the balance appears and how it is consumed.")
step(doc, 7, "Save and map", "Copy the exact Product ID into the admin package Store Products dialog with provider App Store and type Consumable.")

h(doc, "Lifetime courses", 1)
para(doc, "Follow the same path, but select Non-Consumable. Use one product per course and entitlement meaning. In admin, select Lifetime; the derived product type must be Non-consumable, Duration must be empty, and Base Plan ID and Offer ID are not applicable.")

h(doc, "Limited day courses", 1)
para(doc, "These are non-renewing subscriptions. Apple places their management under Monetization, Subscriptions, Non-Renewing Subscriptions, Manage.")
step(doc, 1, "Create a non-renewing subscription", "Enter the Reference Name and Product ID, then set the price and localizations.")
step(doc, 2, "Match the duration", "The access duration shown in the Display Name, Description, paywall, admin Duration days, and review notes must agree. Apple does not renew this product automatically; Italir Pothe stores the expiry.")
step(doc, 3, "Map the product", "In the course Store Products dialog choose App Store, set Access Type to Time limited, set Duration days, and copy the Product ID. The admin representation says subscription; the signed Apple transaction must be NON_RENEWING_SUBSCRIPTION.")
step(doc, 4, "Do not add a base plan", "Base Plan ID is a Google Play field and must remain empty for App Store mappings.")

title(doc, "Monthly Streak Subscription", page_break=True)
h(doc, "Create the subscription group", 1)
para(doc, "Create one group named Streak Protection. Put the monthly unlimited streak plan in this group. Users can hold only one active subscription in a group, which prevents duplicate streak plans.")
step(doc, 1, "Open Subscriptions", "Under Monetization, select Subscriptions and create the Streak Protection group if it does not already exist.")
step(doc, 2, "Create the product", "Inside the group, create a subscription with the final Product ID.")
step(doc, 3, "Set duration", "Choose one month. Set the standard App Store price and storefront availability.")
step(doc, 4, "Add localizations", "Add the group localization and subscription localization. Use language that says the plan automatically renews until cancelled.")
step(doc, 5, "Complete review information", "Upload the paywall screenshot. The screen should show the localized price, monthly duration, automatic renewal, restore action, privacy policy, and terms link.")
step(doc, 6, "Map in admin", "Create a streak_freeze package with monthly billing, monthly_unlimited protection, and normally 30 protectionDurationDays. Add an App Store Store Product with type Subscription, the same Product ID, no Base Plan ID, and no Offer ID for the standard plan.")

h(doc, "Renewal behavior", 2)
bullet(doc, "The App Store transaction expiry is the authority for active protection.")
bullet(doc, "The backend accepts subscription lifecycle registration only when billingModel is monthly and streakProtectionMode is monthly_unlimited.")
bullet(doc, "Cancellation stops future renewal but does not normally remove access before the paid period ends.")
bullet(doc, "Refund, revocation, expiration, billing retry, and renewal events are processed through App Store Server Notifications Version 2.")

h(doc, "Submission rule", 2)
para(doc, "Apple requires the first product of each purchase type to be submitted with a new app version. The first auto-renewable subscription must be submitted with its subscription group and a new app version. Later products of an already approved type can be submitted separately, subject to the current App Store Connect workflow.")

title(doc, "Admin Portal Mapping", page_break=True)
h(doc, "Courses", 1)
para(doc, "Open Admin, Course Directory, edit the course, then open Store Products and add a mapping.")
table(doc, ["Admin field", "Lifetime", "Limited day"], [
    ("Provider", "App Store", "App Store"),
    ("Product ID", "Exact non-consumable ID", "Exact non-renewing subscription ID"),
    ("Access Type", "Lifetime", "Time limited"),
    ("Duration days", "Empty", "Exact access duration"),
    ("Derived Product Type", "Non-consumable", "Subscription"),
    ("Base Plan ID", "Empty", "Empty"),
    ("Offer ID", "Empty", "Empty in current course purchase flow"),
    ("Active Mapping", "On", "On"),
], [1.65, 2.75, 2.8], 9.1)

h(doc, "Shop packages", 1)
para(doc, "Open Admin, Package Store, choose a package, then open Store Products.")
table(doc, ["Package", "Required package values", "App Store product mapping"], [
    ("AI bundle", "ai_bundle; one_time; voiceMinutes and textTokens greater than zero", "Consumable; exact Product ID; no base plan"),
    ("CV credits", "cv_credit; one_time; cvCreditCount greater than zero", "Consumable; exact Product ID; no base plan"),
    ("Finite freezes", "streak_freeze; finite; one_time; freezeCount greater than zero", "Consumable; exact Product ID; no base plan"),
    ("Monthly unlimited streak", "streak_freeze; monthly_unlimited; monthly; duration normally 30 days", "Subscription; exact Product ID; no base plan"),
], [1.45, 3.2, 2.55], 8.9)

h(doc, "Mapping invariant", 2)
para(doc, "Never change a product mapping to represent a different entitlement after orders exist. Deactivate the old mapping and create a new App Store product and a new mapping. Product IDs are transaction keys, not display labels.")

title(doc, "Coupon Architecture", page_break=True)
h(doc, "The central rule", 1)
para(doc, "A backend coupon percentage does not change Apple’s checkout price. The coupon must resolve to an App Store product or Apple subscription offer whose price was configured in App Store Connect. The store remains the charge authority.")

coupon_rows = [
    ("AI, CV, finite freezes", "Separate discounted Consumable", "coupon_ plus the regular Product ID", "Create the second consumable at the discounted App Store price"),
    ("Lifetime course", "Separate discounted Non Consumable", "coupon_ plus the regular Product ID", "Create the second non-consumable at the discounted price"),
    ("Limited day course", "Separate discounted Non Renewing Subscription", "coupon_ plus the regular Product ID", "Create a separate non-renewing product with the same access duration"),
    ("Monthly streak", "Promotional Offer", "Same Product ID plus Promotional Offer ID", "Best for eligible existing or expired subscribers; requires a server signature"),
    ("Monthly streak", "Offer Code", "Same Product ID plus offer reference or identifier", "Best for a code the customer redeems through Apple’s redemption sheet"),
]
table(doc, ["Product", "Apple discount mechanism", "Backend mapping", "Configuration"], coupon_rows, [1.25, 1.65, 1.75, 2.55], 8.65)

h(doc, "Separate discounted products", 2)
para(doc, "For non-subscription native offers, the backend requires the discounted provider Product ID to be exactly coupon_ followed by the regular provider Product ID. Example: regular com.shafacode.italirpothe.cv.credits.5.v1 becomes coupon_com.shafacode.italirpothe.cv.credits.5.v1. Both mappings remain active. Create both products in App Store Connect and set the discounted product’s price there.")
para(doc, "A separate Apple product is not a native coupon. Apple sees it as a distinct IAP. Keep it hidden behind validated coupon checkout in the app, document the behavior for App Review, and verify that duplicate entitlement protection remains correct.")

h(doc, "Monthly promotional offer", 2)
step(doc, 1, "Create the offer", "Open the monthly streak subscription, add a Promotional Offer, and record the Promotional Offer Identifier.")
step(doc, 2, "Set eligibility and price", "Promotional offers are for existing and previously subscribed users. Choose pay as you go, pay up front, or free and the duration. Apple allows up to 10 active promotional offers per subscription.")
step(doc, 3, "Map the coupon", "In Influencer Hub provider mappings, choose App Store, use the same regular and discounted Product ID, choose Promotional Offer, and enter the exact Promotional Offer ID. Base-plan fields remain empty.")
step(doc, 4, "Sign on the server", "The backend creates a time-sensitive signature using the In-App Purchase key, Product ID, Offer ID, order UUID as appAccountToken, nonce, and timestamp. Never place the private key in the app.")

title(doc, "Coupon Readiness by Flow", page_break=True)
h(doc, "Monthly offer code", 2)
para(doc, "Use Apple Offer Codes for a typed campaign code. Configure customer eligibility for new, existing, or expired subscribers, countries or regions, duration, price mode, and redemption limits. In the mapping choose Offer Code and enter the exact Apple offer reference expected by StoreKit and backend. The app opens Apple’s code redemption sheet; Apple decides eligibility and final price.")
table(doc, ["Flow", "Current readiness", "Important limitation"], [
    ("Consumable packages with coupon_ product", "Supported", "Requires a separately created discounted App Store product and exact mapping"),
    ("Lifetime course with coupon_ product", "Supported by backend mapping pattern", "Test end to end because Apple treats regular and discounted products as separate non-consumables"),
    ("Limited course with coupon_ product", "Supported by non-renewing product pattern", "Duration must match; no Apple promotional offer applies"),
    ("Monthly streak promotional offer", "Implemented", "Requires same Product ID, exact Offer ID, In-App Purchase key, and eligible subscriber"),
    ("Monthly streak offer code", "Implemented", "Redemption occurs in Apple UI; validate offer identifier returned in transaction"),
    ("Course auto-renewable promotional offer", "Not applicable to the current course model", "Courses use lifetime or non-renewing access; the course client does not launch signed promotional offers"),
], [1.8, 2.0, 3.4], 8.9)

h(doc, "Coupon consistency checks", 2)
bullet(doc, "Coupon code percentage and the App Store configured price must represent the same campaign. Do not calculate or promise a final local-currency amount from the backend quote.")
bullet(doc, "Campaign dates in admin should match the Apple offer or discounted-product availability window.")
bullet(doc, "For promotional offers, discountedProviderProductId equals regularProviderProductId, providerOfferId is required, appStoreOfferType is promotional_offer, and both base-plan IDs are empty.")
bullet(doc, "For offer codes, choose appStoreOfferType offer_code and verify the identifier stored by the backend matches the identifier Apple returns in the signed transaction.")

h(doc, "Potential naming trap", 2)
para(doc, "The admin form labels the offer-code value as an Apple Offer Code Reference Name. Before production, confirm on a sandbox transaction that this value equals StoreKit’s transaction offerIdentifier. If the current App Store Connect screen distinguishes a marketing reference name from a machine identifier, send a screenshot so the mapping label and stored value can be verified precisely.")

title(doc, "Server Configuration and Notifications", page_break=True)
h(doc, "Backend values required for real verification", 1)
table(doc, ["Environment value", "Purpose"], [
    ("APP_STORE_VERIFICATION_MODE", "Set to real in production so the backend verifies signed Apple transactions."),
    ("APP_STORE_BUNDLE_ID", "Must be com.shafacode.italirpothe."),
    ("APP_STORE_APP_APPLE_ID", "Numeric Apple ID of the App Store Connect app."),
    ("APP_STORE_ISSUER_ID", "Issuer ID associated with App Store Connect API access."),
    ("APP_STORE_KEY_ID", "Key ID for the In-App Purchase key used for server API and promotional-offer signing."),
    ("APP_STORE_PRIVATE_KEY", "Private .p8 key material stored only in the backend secret manager."),
    ("APP_STORE_ROOT_CA_CERTS_BASE64", "Comma-separated base64 DER Apple root certificates used to verify signed data."),
    ("APP_STORE_ENABLE_ONLINE_CERTIFICATE_CHECKS", "Backend control for online certificate checking."),
    ("APP_STORE_NOTIFICATION_ENCRYPTION_KEY", "32-byte key used to encrypt retained notification payloads."),
    ("APP_STORE_NOTIFICATION_MAX_ATTEMPTS", "Retry limit for processing server notifications."),
], [2.55, 4.65], 9.0)
para(doc, "The repository’s sample environment file does not currently list this App Store block. Add documented placeholders to deployment configuration without committing real key material.", italic=True)

h(doc, "App Store Server Notifications", 1)
para(doc, "In App Store Connect, open the app, App Information, and locate App Store Server Notifications. Configure both Production and Sandbox URLs using Version 2.")
para(doc, "Backend route  POST /billing/app-store/notifications")
para(doc, "Use the public HTTPS API origin followed by that route. Example only: https://api.example.com/billing/app-store/notifications. Do not copy the example domain.")
bullet(doc, "If no Sandbox URL is set, Apple sends sandbox notifications to the Production URL. Setting both separately is clearer for testing.")
bullet(doc, "Confirm the endpoint returns HTTP 200 after accepting the signedPayload and logs notificationUUID without logging private keys or full sensitive payloads.")
bullet(doc, "Exercise renewal, expiration, cancellation, billing retry, refund, revocation, and refund-reversal events.")

h(doc, "Key creation", 2)
para(doc, "In App Store Connect, open Users and Access, Integrations, Keys, In-App Purchase, and generate an In-App Purchase key. Download it once, store it securely, and capture the Key ID and Issuer ID. Revoke and replace it immediately if it is lost or exposed.")

title(doc, "Test and Review Checklist", page_break=True)
h(doc, "Catalog validation", 1)
for x in [
    "Every App Store Product ID exactly matches the active admin mapping.",
    "Every product returns a localized title and price on a physical iPhone signed into Sandbox.",
    "Limited-course duration text matches admin durationDays and the backend expiry result.",
    "Only monthly unlimited streak protection is configured as an auto-renewable subscription.",
    "No App Store mapping contains a Google Play Base Plan ID.",
    "Each standard and coupon_ product has its own intended App Store price and availability.",
]:
    bullet(doc, "☐ " + x)

h(doc, "Transaction tests", 1)
for x in [
    "Buy each consumable twice and confirm the wallet increments twice.",
    "Spend AI, CV, and streak balances and confirm each balance decrements correctly.",
    "Buy a lifetime course, delete and reinstall the app, use Restore Purchases, and confirm access returns.",
    "Buy a limited-day course and verify the exact expiresAt value. Confirm it does not auto-renew.",
    "Buy the monthly streak subscription, cancel renewal, and confirm protection remains until the current period expires.",
    "Restore the monthly subscription on a new installation and verify it links to the correct Italir Pothe account.",
    "Attempt account mismatch and Product ID mismatch transactions and confirm the backend rejects them.",
    "Trigger refund and revocation in sandbox and confirm entitlements are reversed without creating negative balances.",
]:
    bullet(doc, "☐ " + x)

h(doc, "Coupon tests", 1)
for x in [
    "Invalid coupon does not open Apple checkout.",
    "Valid consumable or course coupon selects exactly coupon_ plus the regular Product ID.",
    "Displayed purchase price comes from StoreKit for the selected product or offer.",
    "Promotional-offer signature uses the order UUID as appAccountToken and succeeds only for eligible users.",
    "Offer-code redemption opens Apple’s sheet and the verified transaction contains the expected offer identifier.",
    "Expired, paused, wrong-product, and wrong-provider coupons are rejected.",
]:
    bullet(doc, "☐ " + x)

h(doc, "App Review package", 1)
for x in [
    "Agreement, tax, and banking setup is active in App Store Connect.",
    "Each item has at least one localization, price, availability, screenshot, and review notes.",
    "The review account can reach all paywalls without external assistance.",
    "Subscription paywall includes price and duration, renewal statement, restore control, Terms of Use, and Privacy Policy.",
    "First products of each Apple type are attached to the required app version submission.",
]:
    bullet(doc, "☐ " + x)

title(doc, "Screenshot Reconciliation Worksheet", page_break=True)
para(doc, "To convert this setup guide into an exact audit of the current App Store Connect account, provide the screenshots below. Screenshots may redact financial totals, personal names, and secret keys; Product IDs, types, durations, statuses, and offer identifiers must remain visible.")
h(doc, "Screenshots needed", 1)
for x in [
    "Monetization, In-App Purchases list, including all consumables and non-consumables.",
    "Monetization, Subscriptions overview, including Non-Renewing Subscriptions and subscription groups.",
    "Streak Protection subscription group and the monthly product detail page.",
    "Each Promotional Offer detail or list showing its identifier, duration, type, and price.",
    "Each Offer Code configuration showing the offer reference or identifier and customer eligibility.",
    "App Information, App Store Server Notifications section with hostnames and Version 2 visible; redact private query strings if any.",
    "One product detail page of each type showing localization, price, availability, and review information status.",
]:
    bullet(doc, "☐ " + x)

h(doc, "Exact mapping inventory", 1)
table(doc, ["Feature or package", "ASC type", "Product ID", "Offer ID or code", "Admin mapping status"], [
    ("", "", "", "", ""),
    ("", "", "", "", ""),
    ("", "", "", "", ""),
    ("", "", "", "", ""),
    ("", "", "", "", ""),
], [1.55, 1.15, 2.2, 1.3, 1.0], 8.5)

h(doc, "Sign off", 2)
table(doc, ["Check", "Owner", "Date", "Result"], [
    ("ASC products match this mapping", "", "", ""),
    ("Admin mappings match ASC", "", "", ""),
    ("Sandbox purchase suite passes", "", "", ""),
    ("Server notifications pass", "", "", ""),
    ("App Review metadata complete", "", "", ""),
], [3.2, 1.3, 1.2, 1.5], 9.0)

title(doc, "Authoritative References", page_break=True)
para(doc, "Apple documentation checked for this guide on 7 September 2026. App Store Connect labels can change; Apple’s current help pages remain the final authority for screen placement and field names.")
refs = [
    ("Apple overview of In-App Purchase types", "https://developer.apple.com/help/app-store-connect/configure-in-app-purchase-settings/overview-for-configuring-in-app-purchases"),
    ("Create consumable or non-consumable products", "https://developer.apple.com/help/app-store-connect/manage-in-app-purchases/create-consumable-or-non-consumable-in-app-purchases/"),
    ("Create non-renewing subscriptions", "https://developer.apple.com/help/app-store-connect/manage-in-app-purchases/create-non-renewing-subscriptions/"),
    ("Offer auto-renewable subscriptions", "https://developer.apple.com/help/app-store-connect/manage-subscriptions/offer-auto-renewable-subscriptions/"),
    ("In-App Purchase field limits and review information", "https://developer.apple.com/help/app-store-connect/reference/in-app-purchases-and-subscriptions/in-app-purchase-information/"),
    ("Set up promotional offers", "https://developer.apple.com/help/app-store-connect/manage-subscriptions/set-up-promotional-offers-for-auto-renewable-subscriptions/"),
    ("Set up subscription offer codes", "https://developer.apple.com/help/app-store-connect/manage-subscriptions/set-up-subscription-offer-codes/"),
    ("Generate In-App Purchase keys", "https://developer.apple.com/help/app-store-connect/configure-in-app-purchase-settings/generate-keys-for-in-app-purchases/"),
    ("Enter App Store Server Notification URLs", "https://developer.apple.com/help/app-store-connect/configure-in-app-purchase-settings/enter-server-urls-for-app-store-server-notifications"),
    ("Submit an In-App Purchase", "https://developer.apple.com/help/app-store-connect/manage-submissions-to-app-review/submit-an-in-app-purchase/"),
    ("Testing In-App Purchases", "https://developer.apple.com/documentation/storekit/testing-at-all-stages-of-development-with-xcode-and-the-sandbox"),
]
for name, url in refs:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(5)
    add_hyperlink(p, name, url)

h(doc, "Repository evidence used", 1)
table(doc, ["Repository area", "Verified behavior"], [
    ("Flutter shop billing", "Queries StoreKit by Product ID, uses buyConsumable for consumables, launches signed promotional offers, opens Apple code redemption, and sends the order UUID as appAccountToken."),
    ("Course commerce", "Lifetime course requires non-consumable; limited course expects Apple NON_RENEWING_SUBSCRIPTION and stores an expiry."),
    ("Package store", "One-time AI, CV, and finite streak packages require consumables; monthly packages require subscription products."),
    ("Subscription lifecycle", "Only monthly unlimited streak-protection packages are accepted for App Store subscription lifecycle processing."),
    ("Influencer coupon mapping", "Separate products require coupon_ prefix; native Apple subscription offers use the same Product ID, an Offer ID, an offer type, and no base-plan fields."),
], [1.65, 5.55], 8.8)

doc.core_properties.title = "App Store Purchase Mapping Guide"
doc.core_properties.subject = "Italir Pothe App Store Connect product and coupon mapping"
doc.core_properties.author = "OpenAI Codex for Italir Pothe"
doc.core_properties.keywords = "App Store Connect, StoreKit, in-app purchase, subscription, coupon, Italir Pothe"
doc.save(OUT)
print(OUT.resolve())
